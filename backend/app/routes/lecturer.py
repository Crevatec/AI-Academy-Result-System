"""
backend/app/routes/lecturer.py  (FIXED)
----------------------------------------
Fix: enter_scores() now reads total_score_<id> field names
     matching the updated enter_scores.html template.
"""

import os
import uuid
import io
from flask import (Blueprint, render_template, redirect, url_for,
                   flash, request, current_app, send_file)
from flask_login import current_user

from app import db
from app.models.user import Lecturer, User
from app.models.course import Course
from app.models.result import Result, ResultDetail, AcademicSession
from app.models.user import Student
from app.services.grading import compute_grade
from app.services.email_service import send_hod_submission_notification
from app.utils.decorators import lecturer_required
from app.utils.excel_parser import parse_score_sheet, allowed_file

lecturer_bp = Blueprint("lecturer", __name__)


def get_current_lecturer() -> Lecturer:
    return Lecturer.query.filter_by(user_id=current_user.id).first()


def get_current_session() -> AcademicSession:
    return AcademicSession.query.filter_by(is_current=True).first()


# ── Dashboard ─────────────────────────────────────────────────────────────────

@lecturer_bp.route("/dashboard")
@lecturer_required
def dashboard():
    lecturer = get_current_lecturer()
    current_session = get_current_session()

    course_summaries = []
    for course in lecturer.courses:
        if not current_session:
            break

        results = Result.query.filter_by(
            course_id=course.id,
            session_id=current_session.id,
            submitted_by=current_user.id,
        ).all()

        submitted_count = len(results)
        approved_count  = sum(1 for r in results if r.status == "approved")
        pending_count   = sum(1 for r in results if r.status == "pending")
        draft_count     = sum(1 for r in results if r.status == "draft")
        rejected_count  = sum(1 for r in results if r.status == "rejected")

        if approved_count == submitted_count and submitted_count > 0:
            status = "approved"
        elif pending_count > 0:
            status = "pending"
        elif draft_count > 0:
            status = "draft"
        elif rejected_count > 0:
            status = "rejected"
        else:
            status = "not_started"

        course_summaries.append({
            "course":          course,
            "submitted_count": submitted_count,
            "status":          status,
            "rejected_count":  rejected_count,
        })

    return render_template(
        "lecturer/dashboard.html",
        lecturer=lecturer,
        course_summaries=course_summaries,
        current_session=current_session,
    )


# ── Enter Scores (FIXED) ──────────────────────────────────────────────────────

@lecturer_bp.route("/course/<int:course_id>/enter-scores", methods=["GET", "POST"])
@lecturer_required
def enter_scores(course_id: int):
    """
    Single total score entry per student (0-100).
    The template sends field name: total_score_<student_id>
    Grade and GP are computed automatically.
    CA (30%) + Exam (70%) split is stored in ResultDetail for transcript use.
    """
    lecturer = get_current_lecturer()
    current_session = get_current_session()

    if not current_session:
        flash("No active academic session. Contact Admin.", "warning")
        return redirect(url_for("lecturer.dashboard"))

    course = Course.query.get_or_404(course_id)
    if course not in lecturer.courses:
        flash("You are not assigned to this course.", "danger")
        return redirect(url_for("lecturer.dashboard"))

    students = (
        Student.query
        .filter_by(department_id=course.department_id, level=course.level)
        .join(Student.user)
        .all()
    )

    existing = {
        r.student_id: r
        for r in Result.query.filter_by(
            course_id=course_id,
            session_id=current_session.id,
            submitted_by=current_user.id,
        ).all()
    }

    if request.method == "POST":
        saved   = 0
        errors  = 0

        for student in students:
            # ── Read the total score field ─────────────────────────────────
            score_str = request.form.get(
                f"total_score_{student.id}", ""
            ).strip()

            if not score_str:
                continue  # student left blank — skip

            try:
                total = float(score_str)
            except ValueError:
                flash(
                    f"Invalid score '{score_str}' for "
                    f"{student.user.full_name}. Must be a number.",
                    "warning"
                )
                errors += 1
                continue

            if not (0 <= total <= 100):
                flash(
                    f"Score {total} for {student.user.full_name} "
                    f"must be between 0 and 100.",
                    "warning"
                )
                errors += 1
                continue

            scale = course.department.grading_scale
            grade, grade_point = compute_grade(total, scale)

            # Split into CA 30% + Exam 70% for detail record
            ca_score   = round(total * 0.30, 1)
            exam_score = round(total - ca_score, 1)

            if student.id in existing:
                result = existing[student.id]
                if result.status == "draft":
                    result.score       = total
                    result.grade       = grade
                    result.grade_point = grade_point
                    if result.detail:
                        result.detail.ca_score   = ca_score
                        result.detail.exam_score = exam_score
                    else:
                        detail = ResultDetail(
                            result_id       = result.id,
                            ca_score        = ca_score,
                            ca_max          = 30.0,
                            practical_score = None,
                            practical_max   = 0.0,
                            exam_score      = exam_score,
                            exam_max        = 70.0,
                        )
                        db.session.add(detail)
                    saved += 1
                # pending/approved: skip without error
            else:
                result = Result(
                    student_id   = student.id,
                    course_id    = course_id,
                    session_id   = current_session.id,
                    submitted_by = current_user.id,
                    score        = total,
                    grade        = grade,
                    grade_point  = grade_point,
                    status       = "draft",
                )
                db.session.add(result)
                db.session.flush()

                detail = ResultDetail(
                    result_id       = result.id,
                    ca_score        = ca_score,
                    ca_max          = 30.0,
                    practical_score = None,
                    practical_max   = 0.0,
                    exam_score      = exam_score,
                    exam_max        = 70.0,
                )
                db.session.add(detail)
                saved += 1

        db.session.commit()

        if saved > 0:
            flash(
                f"✓ Scores saved for {saved} student(s). Status: Draft.",
                "success"
            )
        elif errors == 0:
            flash(
                "No scores entered. Type scores in the boxes then click "
                "'Save as Draft'.",
                "warning"
            )

        return redirect(url_for("lecturer.enter_scores", course_id=course_id))

    return render_template(
        "lecturer/enter_scores.html",
        course=course,
        students=students,
        existing=existing,
        current_session=current_session,
    )


# ── View Submitted Results ────────────────────────────────────────────────────

@lecturer_bp.route("/course/<int:course_id>/view-results")
@lecturer_required
def view_course_results(course_id: int):
    lecturer      = get_current_lecturer()
    current_session = get_current_session()
    course        = Course.query.get_or_404(course_id)

    if course not in lecturer.courses:
        flash("You are not assigned to this course.", "danger")
        return redirect(url_for("lecturer.dashboard"))

    results = Result.query.filter_by(
        course_id    = course_id,
        session_id   = current_session.id,
        submitted_by = current_user.id,
    ).all()

    scores      = [r.score for r in results]
    avg_score   = round(sum(scores) / len(scores), 1) if scores else 0
    pass_count  = sum(1 for r in results if r.grade != "F")
    fail_count  = sum(1 for r in results if r.grade == "F")

    return render_template(
        "lecturer/view_results.html",
        course=course,
        results=results,
        current_session=current_session,
        avg_score=avg_score,
        pass_count=pass_count,
        fail_count=fail_count,
    )


# ── Download Results as Excel ─────────────────────────────────────────────────

@lecturer_bp.route("/course/<int:course_id>/download-results")
@lecturer_required
def download_course_results(course_id: int):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    lecturer        = get_current_lecturer()
    current_session = get_current_session()
    course          = Course.query.get_or_404(course_id)

    if course not in lecturer.courses:
        flash("Unauthorized.", "danger")
        return redirect(url_for("lecturer.dashboard"))

    results = Result.query.filter_by(
        course_id    = course_id,
        session_id   = current_session.id,
        submitted_by = current_user.id,
    ).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{course.code} Results"

    BORDER = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"),  bottom=Side(style="thin")
    )
    HFILL  = PatternFill("solid", fgColor="1A3C6B")
    HFONT  = Font(bold=True, color="FFFFFF")
    CENTER = Alignment(horizontal="center", vertical="center")
    LEFT   = Alignment(horizontal="left",   vertical="center")

    ws.merge_cells("A1:J1")
    ws["A1"].value = (
        f"{course.code} — {course.title}  |  "
        f"{current_session.session_name} Semester {current_session.semester}  |  "
        f"Lecturer: {current_user.full_name}"
    )
    ws["A1"].font      = Font(bold=True, size=12, color="1A3C6B")
    ws["A1"].alignment = CENTER
    ws.row_dimensions[1].height = 25

    headers = [
        "S/N", "Matric No", "Student Name",
        "CA (/30)", "Practical", "Exam (/70)",
        "Total (/100)", "Grade", "GP", "Status"
    ]
    widths = [5, 16, 24, 10, 10, 10, 12, 8, 8, 12]
    for ci, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=2, column=ci, value=h)
        c.font = HFONT; c.fill = HFILL
        c.border = BORDER; c.alignment = CENTER
        ws.column_dimensions[c.column_letter].width = w

    GRADE_COLORS = {
        "A": "C8E6C9", "B": "DCEDC8", "C": "FFF9C4",
        "D": "FFE0B2", "E": "FFCCBC", "F": "FFCDD2"
    }

    for sn, result in enumerate(results, start=1):
        detail = result.detail
        rfill  = PatternFill("solid",
                             fgColor=GRADE_COLORS.get(result.grade, "FFFFFF"))
        row_vals = [
            sn,
            result.student.matric_number,
            result.student.user.full_name,
            detail.ca_score        if detail else "—",
            detail.practical_score if detail else "—",
            detail.exam_score      if detail else "—",
            round(result.score, 1),
            result.grade,
            result.grade_point,
            result.status.capitalize(),
        ]
        for ci, val in enumerate(row_vals, start=1):
            c = ws.cell(row=sn + 2, column=ci, value=val)
            c.border = BORDER; c.fill = rfill
            c.alignment = CENTER if ci != 3 else LEFT

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = (
        f"{course.code}_Results_"
        f"{current_session.session_name.replace('/', '-')}_"
        f"Sem{current_session.semester}.xlsx"
    )
    return send_file(
        buffer,
        mimetype=(
            "application/vnd.openxmlformats-officedocument"
            ".spreadsheetml.sheet"
        ),
        as_attachment=True,
        download_name=filename,
    )


# ── Excel Upload ──────────────────────────────────────────────────────────────

@lecturer_bp.route("/course/<int:course_id>/upload-scores", methods=["POST"])
@lecturer_required
def upload_scores(course_id: int):
    lecturer        = get_current_lecturer()
    current_session = get_current_session()

    if not current_session:
        flash("No active session.", "warning")
        return redirect(url_for("lecturer.dashboard"))

    course = Course.query.get_or_404(course_id)
    if course not in lecturer.courses:
        flash("Unauthorized.", "danger")
        return redirect(url_for("lecturer.dashboard"))

    file = request.files.get("excel_file")
    if not file or not file.filename:
        flash("Please select an Excel file.", "warning")
        return redirect(url_for("lecturer.enter_scores", course_id=course_id))

    if not allowed_file(file.filename, current_app.config["ALLOWED_EXTENSIONS"]):
        flash("Invalid file type. Upload a .xlsx file.", "danger")
        return redirect(url_for("lecturer.enter_scores", course_id=course_id))

    temp_filename = f"{uuid.uuid4().hex}.xlsx"
    upload_folder = current_app.config["UPLOAD_FOLDER"]
    os.makedirs(upload_folder, exist_ok=True)
    temp_path = os.path.join(upload_folder, temp_filename)
    file.save(temp_path)

    records, errors = parse_score_sheet(temp_path)
    os.remove(temp_path)

    for err in errors:
        flash(err, "warning")

    scale     = course.department.grading_scale
    saved     = 0
    not_found = []

    for record in records:
        student = Student.query.filter_by(
            matric_number=record["matric"]
        ).first()
        if not student:
            not_found.append(record["matric"])
            continue

        grade, grade_point = compute_grade(record["score"], scale)
        ca_score   = round(record["score"] * 0.30, 1)
        exam_score = round(record["score"] * 0.70, 1)

        existing = Result.query.filter_by(
            student_id = student.id,
            course_id  = course_id,
            session_id = current_session.id,
        ).first()

        if existing:
            if existing.status == "draft":
                existing.score       = record["score"]
                existing.grade       = grade
                existing.grade_point = grade_point
                if existing.detail:
                    existing.detail.ca_score   = ca_score
                    existing.detail.exam_score = exam_score
        else:
            result = Result(
                student_id   = student.id,
                course_id    = course_id,
                session_id   = current_session.id,
                submitted_by = current_user.id,
                score        = record["score"],
                grade        = grade,
                grade_point  = grade_point,
                status       = "draft",
            )
            db.session.add(result)
            db.session.flush()
            detail = ResultDetail(
                result_id       = result.id,
                ca_score        = ca_score,
                ca_max          = 30.0,
                practical_score = None,
                practical_max   = 0.0,
                exam_score      = exam_score,
                exam_max        = 70.0,
            )
            db.session.add(detail)

        saved += 1

    db.session.commit()
    flash(f"Upload complete. {saved} records saved as Draft.", "success")
    if not_found:
        flash(
            f"Matric numbers not found in system: {', '.join(not_found)}",
            "warning"
        )

    return redirect(url_for("lecturer.enter_scores", course_id=course_id))


# ── Submit to HOD ─────────────────────────────────────────────────────────────

@lecturer_bp.route("/course/<int:course_id>/submit", methods=["POST"])
@lecturer_required
def submit_results(course_id: int):
    lecturer        = get_current_lecturer()
    current_session = get_current_session()
    course          = Course.query.get_or_404(course_id)

    if course not in lecturer.courses:
        flash("Unauthorized.", "danger")
        return redirect(url_for("lecturer.dashboard"))

    drafts = Result.query.filter_by(
        course_id    = course_id,
        session_id   = current_session.id,
        submitted_by = current_user.id,
        status       = "draft",
    ).all()

    if not drafts:
        flash("No draft results to submit for this course.", "info")
        return redirect(url_for("lecturer.dashboard"))

    for result in drafts:
        result.status = "pending"

    db.session.commit()

    hod_uid = course.department.hod_user_id
    if hod_uid:
        hod = User.query.get(hod_uid)
        if hod:
            send_hod_submission_notification(
                hod.email,
                current_user.full_name,
                course.code,
                len(drafts),
            )

    flash(
        f"✓ {len(drafts)} result(s) for {course.code} submitted to HOD.",
        "success"
    )
    return redirect(url_for("lecturer.dashboard"))
